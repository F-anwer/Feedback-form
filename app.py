import streamlit as st
import time
import io
import pandas as pd
from datetime import datetime
from supabase import create_client, Client
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Initialize Supabase Client
SUPABASE_URL = st.secrets["SUPABASE_URL"]
SUPABASE_KEY = st.secrets["SUPABASE_KEY"]
ADMIN_PASSWORD = st.secrets["ADMIN_PASSWORD"]

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

st.set_page_config(page_title="Moellim Course Page Review", layout="wide")

# ==========================================
# CONFIGURATION: SUBMISSION DEADLINE
# ==========================================
SUBMISSION_DEADLINE = datetime(2026, 10, 26, 23, 59, 59)

# Sidebar Role Selector
st.sidebar.title("Navigation")
role = st.sidebar.radio("Select View:", ["Instructor Portal", "Evaluator (Admin) Portal"])

now = datetime.now()
is_deadline_passed = now > SUBMISSION_DEADLINE

st.sidebar.divider()
st.sidebar.subheader("📅 Submission Deadline")
st.sidebar.caption(f"**Deadline:** {SUBMISSION_DEADLINE.strftime('%b %d, %Y at %I:%M %p')}")
if is_deadline_passed:
    st.sidebar.error("🔒 Submission Portal Closed")
else:
    st.sidebar.success("🟢 Submissions Active")


# Dialog Pop-up for Submission Confirmation
@st.dialog("🎉 Submission Received!")
def show_submission_popup(course, tracking_token, timestamp_str):
    st.success("Your course page PDF has been successfully uploaded for QEC review.")
    st.write(f"**Course:** {course}")
    st.write(f"**Tracking Token:** `{tracking_token}`")
    st.write(f"**Submitted On:** {timestamp_str}")
    st.info("You can enter your registered email under 'Track Submission Status' to view reviews and comments.")
    if st.button("Close Window"):
        st.rerun()


# Function to generate Professionally Formatted & Color-Coded Excel Summary Report
def generate_excel_report(data):
    if not data:
        return None
    
    df = pd.DataFrame(data)
    
    export_df = pd.DataFrame({
        "Submission ID": df.get("id", ""),
        "Instructor Name": df.get("instructor_name", ""),
        "Instructor Email": df.get("email", ""),
        "Course Title & Code": df.get("course_title", ""),
        "Submission Status": df.get("status", "Pending Review"),
        "Submitted At": df.get("created_at", "").apply(lambda x: str(x)[:10] if str(x) != "nan" and x else ""),
        "Reviewer Feedback": df.get("feedback", ""),
        "PDF Document Link": df.get("pdf_url", "")
    })
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="QEC Submissions Summary")
        workbook = writer.book
        worksheet = writer.sheets["QEC Submissions Summary"]

        # --- STYLES DEFINITION ---
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        row_alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        row_white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        # Status Color Fills & Fonts
        status_styles = {
            "Approved": {
                "fill": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
                "font": Font(name="Calibri", size=11, bold=True, color="155724")
            },
            "Changes Requested": {
                "fill": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
                "font": Font(name="Calibri", size=11, bold=True, color="856404")
            },
            "Pending Review": {
                "fill": PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid"),
                "font": Font(name="Calibri", size=11, bold=True, color="004085")
            },
            "Rejected": {
                "fill": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
                "font": Font(name="Calibri", size=11, bold=True, color="721C24")
            }
        }

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # --- APPLY HEADER STYLES ---
        for col_num, col_name in enumerate(export_df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            worksheet.row_dimensions[1].height = 25

        # --- APPLY DATA ROW STYLES ---
        status_col_idx = export_df.columns.get_loc("Submission Status") + 1

        for row_num in range(2, len(export_df) + 2):
            worksheet.row_dimensions[row_num].height = 22
            is_even = row_num % 2 == 0
            base_fill = row_alt_fill if is_even else row_white_fill

            for col_num in range(1, len(export_df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.border = thin_border
                cell.alignment = align_left
                cell.font = Font(name="Calibri", size=10)
                cell.fill = base_fill

                # Color Code the "Submission Status" column
                if col_num == status_col_idx:
                    status_val = str(cell.value)
                    style_cfg = status_styles.get(status_val, status_styles["Pending Review"])
                    cell.fill = style_cfg["fill"]
                    cell.font = style_cfg["font"]
                    cell.alignment = align_center

        # --- AUTO-FIT COLUMN WIDTHS ---
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    return output.getvalue()


# ==========================================
# 1. INSTRUCTOR PORTAL (PUBLIC)
# ==========================================
if role == "Instructor Portal":
    st.title("Moellim Course Page Portal")
    
    tab_submit, tab_track = st.tabs(["Submit Course Page", "Track Submission Status"])

    # --- SUBMIT TAB ---
    with tab_submit:
        st.header("Submit Course Page for Review")
        
        if is_deadline_passed:
            st.error(f"🔒 The submission portal closed on **{SUBMISSION_DEADLINE.strftime('%B %d, %Y at %I:%M %p')}**. New uploads are no longer accepted.")
            st.info("If you require an extension, please contact the QEC administration team.")
        else:
            with st.form("faculty_submission", clear_on_submit=True):
                course = st.text_input("Course Title & Code *")
                instructor = st.text_input("Instructor Name *")
                email = st.text_input("Instructor Email *")
                pdf = st.file_uploader("Upload Moellim Page PDF *", type=["pdf"])

                submitted = st.form_submit_button("Submit")
                if submitted:
                    if not course or not instructor or not email or not pdf:
                        st.error("Please fill in all required fields and upload your PDF.")
                    elif pdf.size > 25 * 1024 * 1024:
                        st.error("File size exceeds 25MB limit.")
                    else:
                        clean_instructor = "".join(c if c.isalnum() else "_" for c in instructor.strip())
                        clean_course = "".join(c if c.isalnum() else "_" for c in course.strip())
                        
                        timestamp_int = int(time.time())
                        tracking_token = f"QEC-{timestamp_int}"
                        timestamp_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        
                        file_path = f"{clean_instructor}/{clean_course}/{timestamp_int}_submission.pdf"
                        file_bytes = pdf.read()
                        
                        supabase.storage.from_("moellim-pdfs").upload(
                            file_path, 
                            file_bytes, 
                            {"content-type": "application/pdf", "upsert": "true"}
                        )
                        
                        public_url = supabase.storage.from_("moellim-pdfs").get_public_url(file_path)

                        record = {
                            "course_title": course.strip(),
                            "instructor_name": instructor.strip(),
                            "email": email.strip().lower(),
                            "pdf_url": public_url,
                            "status": "Pending Review",
                            "feedback": "Awaiting initial review"
                        }
                        supabase.table("submissions").insert(record).execute()
                        show_submission_popup(course, tracking_token, timestamp_str)

    # --- TRACK STATUS TAB ---
    with tab_track:
        st.header("Track Submission Status")
        search_email = st.text_input("Enter your registered email address:")
        
        if search_email:
            clean_email = search_email.strip().lower()
            response = supabase.table("submissions").select("*").eq("email", clean_email).order("created_at", desc=True).execute()
            data = response.data
            
            if data:
                instructor_name = data[0].get("instructor_name", "Instructor")
                
                st.markdown(f"### 👤 Profile: `{clean_email}`")
                st.caption(f"Showing **{len(data)}** total submission(s) for **{instructor_name}**")
                st.divider()

                course_groups = {}
                for item in data:
                    c_title = item.get("course_title", "Uncategorized Course")
                    course_groups.setdefault(c_title, []).append(item)

                for c_title, submissions in course_groups.items():
                    st.subheader(f"📖 {c_title}")
                    
                    for sub in submissions:
                        status = sub.get("status", "Pending Review")
                        
                        if status == "Approved":
                            badge = "✅ Approved"
                        elif status == "Changes Requested":
                            badge = "⚠️ Changes Requested"
                        elif status == "Rejected":
                            badge = "❌ Rejected"
                        else:
                            badge = "⏳ Pending Review"

                        with st.expander(f"Submission #{sub['id']} — {badge}"):
                            st.write(f"**Submitted Date:** {sub['created_at'][:10]}")
                            st.write(f"**Current Status:** {status}")
                            st.write(f"**Reviewer Feedback:** {sub['feedback']}")
                            if sub.get('pdf_url'):
                                st.markdown(f"[📄 View/Download Uploaded PDF]({sub['pdf_url']})")
                    
                    st.write("") 
            else:
                st.info("No course submissions found associated with this email address.")

# ==========================================
# 2. EVALUATOR / ADMIN PORTAL (PROTECTED)
# ==========================================
elif role == "Evaluator (Admin) Portal":
    st.title("Moellim Page Evaluator Portal")

    if "authenticated" not in st.session_state:
        st.session_state.authenticated = False

    if not st.session_state.authenticated:
        pwd_input = st.text_input("Enter Admin Password:", type="password")
        if st.button("Login"):
            if pwd_input == ADMIN_PASSWORD:
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("Incorrect Password.")
    else:
        if st.sidebar.button("Logout Admin"):
            st.session_state.authenticated = False
            st.rerun()

        st.subheader("Manage & Review Submissions")

        res = supabase.table("submissions").select("*").order("created_at", desc=True).execute()
        rows = res.data

        col_title, col_export = st.columns([3, 1])
        with col_export:
            if rows:
                excel_bytes = generate_excel_report(rows)
                st.download_button(
                    label="📊 Export Formatted Excel Sheet",
                    data=excel_bytes,
                    file_name=f"QEC_Submissions_Report_{datetime.now().strftime('%Y%m%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

        if rows:
            email_groups = {}
            for row in rows:
                instructor_email = row.get("email", "no-email@provided.com").strip().lower()
                email_groups.setdefault(instructor_email, []).append(row)

            st.write(f"Total Active Instructors: **{len(email_groups)}** | Total Submissions: **{len(rows)}**")
            st.divider()

            status_options = ["Pending Review", "Approved", "Changes Requested", "Rejected"]

            for email_addr, submissions in email_groups.items():
                instructor_name = submissions[0].get("instructor_name", "Instructor")
                
                with st.expander(f"📧 {email_addr} — ({instructor_name}) | [{len(submissions)} Submission(s)]"):
                    for row in submissions:
                        st.markdown(f"#### 📖 Course: `{row['course_title']}`")
                        st.write(f"**Submitted Date:** {row['created_at'][:10]}")
                        st.write(f"**Current Status:** `{row.get('status', 'Pending Review')}`")
                        st.markdown(f"[📄 Download/View PDF]({row['pdf_url']})")

                        current_status = row.get("status", "Pending Review")
                        default_index = status_options.index(current_status) if current_status in status_options else 0

                        with st.form(f"eval_form_{row['id']}"):
                            new_status = st.selectbox("Update Status", status_options, index=default_index)
                            new_feedback = st.text_area("Review Comments / Feedback", value=row.get("feedback", ""))
                            
                            saved = st.form_submit_button("Save & Update")
                            if saved:
                                supabase.table("submissions").update({
                                    "status": new_status,
                                    "feedback": new_feedback
                                }).eq("id", row["id"]).execute()
                                st.success(f"Updated status for {row['course_title']}!")
                                st.rerun()

                        if st.button(f"🗑️ Delete Submission #{row['id']}", key=f"del_btn_{row['id']}"):
                            supabase.table("submissions").delete().eq("id", row["id"]).execute()
                            st.success(f"Successfully deleted submission #{row['id']}!")
                            st.rerun()

                        st.markdown("---")
        else:
            st.info("No course page submissions found in database.")